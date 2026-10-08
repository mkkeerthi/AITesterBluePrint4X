"""Test case repositories (CSV / XLSX): one row = one chunk.

A test case is already an atomic unit of meaning, so it is never split and
never merged (the chapter 11 lesson). What this chunker does instead:

* maps whatever the team calls its columns onto canonical fields
* drops columns that are always empty, and columns that duplicate another
  column row for row (this repo's `Steps to Execute` == `TestSteps`)
* writes a compact, labelled text block so every field is searchable
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from ..config import Source
from ..models import Chunk
from ..text import normalize

# canonical field -> header spellings seen in the wild (first match wins)
FIELDS: dict[str, list[str]] = {
    "id": ["scenario tid", "test case id", "testcase id", "tc id", "tc_id", "test id", "tid", "id", "key", "issue key"],
    "title": ["testcase description", "test case description", "test case title", "summary", "title", "test case", "test name", "name"],
    "module": ["module", "category", "feature", "component", "area", "epic"],
    "type": ["test type", "scenario type", "type"],
    "priority": ["priority", "severity"],
    "automated": ["is automated", "automated", "automation status", "automation"],
    "status": ["status", "execution status", "result"],
    "precondition": ["precondition", "preconditions", "pre-condition", "pre conditions", "prerequisites"],
    "data": ["test data", "testdata", "data"],
    "steps": ["teststeps", "test steps", "steps", "steps to execute", "procedure"],
    "expected": ["expected result", "expected results", "expected", "expected outcome"],
    "actual": ["actual result", "actual results", "actual"],
    "labels": ["labels", "tags"],
    "owner": ["executed qa name", "owner", "assignee", "tester", "executed by"],
    "comments": ["misc (comments)", "comments", "notes", "remarks"],
}

LABELS = {
    "precondition": "Precondition",
    "data": "Test data",
    "expected": "Expected result",
    "actual": "Actual result",
    "comments": "Comments",
    "owner": "Executed by",
    "labels": "Labels",
}

_VERIFY = re.compile(r"^\s*verify\s+(?P<module>.+?)\s+-\s+(?P<scenario>.+?)\s*$", re.I)


def _clean_header(h: str) -> str:
    return re.sub(r"\s+", " ", h.replace("﻿", "").strip().strip('"').lower())


def _read_rows(path: Path) -> tuple[list[str], list[list[str]]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        header: list[str] = []
        rows: list[list[str]] = []
        for ws in wb.worksheets:
            sheet_header: list[str] | None = None
            for raw in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in raw]
                if not any(cells):
                    continue
                if sheet_header is None:
                    sheet_header = cells
                    if not header:
                        header = cells
                    continue
                if sheet_header == header:
                    rows.append(cells)
        return header, rows

    # utf-8-sig swallows the BOM that Excel writes (the chapter 11 KeyError)
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader, [])
        rows = [r for r in reader if any(c.strip() for c in r)]
    return header, rows


def _map_columns(header: list[str], rows: list[list[str]]) -> tuple[dict[str, int], list[tuple[str, int]], list[str]]:
    clean = [_clean_header(h) for h in header]
    width = len(header)
    col = lambda i: tuple((r[i] if i < len(r) else "").strip() for r in rows)  # noqa: E731

    dropped: list[str] = []
    keep: list[int] = []
    seen_values: dict[tuple, str] = {}
    for i in range(width):
        values = col(i)
        if not any(values):
            dropped.append(f"{header[i]} (always empty)")
            continue
        if values in seen_values:
            dropped.append(f"{header[i]} (duplicates {seen_values[values]})")
            continue
        seen_values[values] = header[i]
        keep.append(i)

    mapping: dict[str, int] = {}
    for field, names in FIELDS.items():
        for name in names:
            idx = next((i for i in keep if clean[i] == name and i not in mapping.values()), None)
            if idx is not None:
                mapping[field] = idx
                break

    extras = [(header[i].strip(), i) for i in keep if i not in mapping.values()]
    return mapping, extras, dropped


def _steps(raw: str) -> str:
    parts = [p.strip() for p in re.split(r"\s+\|\s+|\n", raw) if p.strip()]
    return "\n".join(parts)


def chunk_testcases(path: Path, source: Source, rel: str) -> list[Chunk]:
    header, rows = _read_rows(path)
    if not header or not rows:
        return []
    mapping, extras, dropped = _map_columns(header, rows)

    chunks: list[Chunk] = []
    for n, row in enumerate(rows, start=1):
        get = lambda f: normalize(row[mapping[f]]) if f in mapping and mapping[f] < len(row) else ""  # noqa: E731
        tc_id, title = get("id"), get("title")
        if not tc_id and not title:
            continue
        tc_id = tc_id or f"ROW-{n}"

        module, scenario = get("module"), ""
        m = _VERIFY.match(title)
        if m:
            module = module or m.group("module")
            scenario = m.group("scenario")
        if not module:
            module = re.sub(r"[-_]?\d+$", "", tc_id).replace("_", " ").title()

        facts = [f"Module: {module}"]
        for f, label in (("priority", "Priority"), ("type", "Type"), ("automated", "Automated"), ("status", "Status")):
            if get(f):
                facts.append(f"{label}: {get(f)}")

        lines = [f"Test case {tc_id}: {title}", " | ".join(facts)]
        for f in ("precondition", "data"):
            if get(f):
                lines.append(f"{LABELS[f]}: {get(f)}")
        if get("steps"):
            lines.append("Steps:\n" + _steps(get("steps")))
        for f in ("expected", "actual", "comments", "owner", "labels"):
            if get(f):
                lines.append(f"{LABELS[f]}: {get(f)}")
        for name, i in extras:
            if i < len(row) and row[i].strip():
                lines.append(f"{name}: {normalize(row[i])}")

        text = "\n".join(lines)
        chunks.append(
            Chunk(
                text=text,
                source_id=source.id,
                source_kind=source.kind,
                file_path=rel,
                title=f"{tc_id} · {title}" if title else tc_id,
                locator=f"row {n}",
                meta={
                    "tc_id": tc_id,
                    "module": module,
                    "scenario": scenario or title,
                    "priority": get("priority"),
                    "automated": get("automated"),
                    "status": get("status"),
                    "test_type": get("type"),
                    "row": n,
                    "dropped_columns": dropped if n == 1 else None,
                },
            )
        )
    return chunks + _summaries(chunks, source, rel, path.name)


def _summaries(rows: list[Chunk], source: Source, rel: str, filename: str) -> list[Chunk]:
    """Inventory chunks for questions about the repository as a whole.

    "Which features have no test cases?" asks about absence, and top-k
    similarity search cannot prove absence: it returns the five most similar
    rows, never "nothing covers heatmaps". One overview chunk plus one chunk
    per module (every scenario name listed) gives gap analysis and RTM
    questions the full inventory to reason over, in a few hundred tokens.
    """
    from collections import Counter, defaultdict

    if not rows:
        return []
    by_mod: dict[str, list[Chunk]] = defaultdict(list)
    for c in rows:
        by_mod[c.meta["module"]].append(c)

    def mix(cs: list[Chunk], key: str) -> str:
        cnt = Counter((c.meta.get(key) or "unknown") for c in cs)
        return ", ".join(f"{k} {v}" for k, v in cnt.most_common())

    auto = sum(1 for c in rows if str(c.meta.get("automated", "")).lower() in {"yes", "y", "true", "automated"})
    overview = [
        f"Test case repository summary ({filename}): {len(rows)} test cases across {len(by_mod)} modules.",
        f"Automated: {auto} ({auto / len(rows):.0%}); manual: {len(rows) - auto}.",
        f"Priority mix: {mix(rows, 'priority')}.",
        f"Execution status: {mix(rows, 'status')}.",
        "Modules (test case count, ID range):",
    ]
    for mod, cs in sorted(by_mod.items(), key=lambda kv: -len(kv[1])):
        ids = [c.meta["tc_id"] for c in cs]
        overview.append(f"- {mod}: {len(cs)} test cases ({ids[0]} to {ids[-1]})")
    overview.append(
        "Only the modules above have test cases. A product feature that maps to none of these modules has no test case in this repository."
    )
    out = [
        Chunk(
            text="\n".join(overview),
            source_id=source.id,
            source_kind=source.kind,
            file_path=rel,
            title=f"Test case repository summary ({len(rows)} cases, {len(by_mod)} modules)",
            locator="summary",
            meta={"summary": "repository", "modules": sorted(by_mod)},
        )
    ]
    for mod, cs in by_mod.items():
        ids = [c.meta["tc_id"] for c in cs]
        auto_m = sum(1 for c in cs if str(c.meta.get("automated", "")).lower() in {"yes", "y", "true", "automated"})
        scen = "; ".join(f"{c.meta['tc_id']} {c.meta.get('scenario', '')}" for c in cs)
        out.append(
            Chunk(
                text=(
                    f"Test case module summary: {mod} ({len(cs)} test cases, {ids[0]} to {ids[-1]}).\n"
                    f"Automated: {auto_m} of {len(cs)}. Priority mix: {mix(cs, 'priority')}.\n"
                    f"Scenarios covered: {scen}"
                ),
                source_id=source.id,
                source_kind=source.kind,
                file_path=rel,
                title=f"Module summary · {mod} ({len(cs)} test cases)",
                locator=f"module {mod}",
                meta={"summary": "module", "module": mod},
            )
        )
    return out
